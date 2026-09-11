"""Bounded CSV/XLSX table extraction retaining physical source locations."""
import csv
import datetime as dt
import io
from pathlib import Path
import re
import zipfile

from .mapping import header_candidates, propose_mapping

MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_EXPANDED_BYTES = 128 * 1024 * 1024
MAX_ROWS = 100_000
MAX_COLUMNS = 256
MAX_CELL_CHARS = 32_768
SUPPORTED_EXTENSIONS = {".csv", ".tsv", ".xlsx"}


def _json_value(value):
    return value.isoformat() if isinstance(value, (dt.datetime, dt.date)) else value


def _is_blank(value):
    return value is None or (isinstance(value, str) and not value.strip())


def _looks_like_header(values):
    nonempty = [str(value).strip() for value in values if not _is_blank(value)]
    if len(nonempty) < 2:
        return False
    recognized = {key for key, matches in header_candidates(nonempty).items() if matches}
    return len(recognized) >= 2 and bool(recognized & {"symbol", "trade_date", "side", "transaction_type", "quantity", "net_amount"})


def _headers(values):
    names, counts = [], {}
    for index, value in enumerate(values, 1):
        name = str(value).strip() if not _is_blank(value) else f"Column {index}"
        counts[name] = counts.get(name, 0) + 1
        names.append(name if counts[name] == 1 else f"{name} [{counts[name]}]")
    return names


def _tables(records, sheet_name="", initial_warnings=None):
    tables, current = [], None
    for row_number, values, metadata in records:
        values = [_json_value(v) for v in values]
        while values and _is_blank(values[-1]):
            values.pop()
        if not values or all(_is_blank(v) for v in values):
            continue
        if len(values) > MAX_COLUMNS:
            raise ValueError(f"A table exceeds the {MAX_COLUMNS}-column import limit.")
        if any(len(str(v)) > MAX_CELL_CHARS for v in values if v is not None):
            raise ValueError("A cell exceeds the maximum supported text length.")
        if _looks_like_header(values):
            headers = _headers(values)
            if current and headers == current["headers"]:
                current["warnings"].append(f"Repeated header at row {row_number} was skipped.")
                continue
            current = {
                "sheet_name": sheet_name, "table_index": len(tables),
                "header_row_index": row_number, "data_start_row_index": row_number + 1,
                "headers": headers, "rows": [], "warnings": list(initial_warnings or []),
            }
            if len(set(str(v) for v in values)) != len(values):
                current["warnings"].append("Duplicate header names were disambiguated; review all charge columns.")
            tables.append(current)
            continue
        if current is None:
            continue
        first = str(next(v for v in values if not _is_blank(v))).strip()
        if re.match(r"(?i)^(?:grand total|sub[ -]?total|total|opening balance|closing balance)(?:\b|:)", first):
            current["warnings"].append(f"Summary row {row_number} was skipped.")
            continue
        if len([v for v in values if not _is_blank(v)]) == 1 and not re.match(r"\d{1,4}[/.-]\d", first):
            current["warnings"].append(f"Note/title row {row_number} was skipped.")
            continue
        row_ref = f"{sheet_name} row {row_number}" if sheet_name else f"row {row_number}"
        if len(values) > len(current["headers"]):
            raise ValueError(f"{row_ref} has more cells than its header; repair the delimiter or table before import.")
        padded = values + [None] * (len(current["headers"]) - len(values))
        row = {"row_ref": row_ref, "values": dict(zip(current["headers"], padded))}
        if metadata.get("warnings"):
            row["warnings"] = metadata["warnings"]
        if metadata.get("formulas"):
            row["formulas"] = {current["headers"][i]: formula for i, formula in metadata["formulas"].items() if i < len(current["headers"])}
        current["rows"].append(row)
    for table in tables:
        suggestion = propose_mapping(table["headers"], table["rows"])
        table.update({key: suggestion[key] for key in ("mapping", "confidence", "detected_type")})
        table["warnings"].extend(suggestion["warnings"])
        table["row_count"] = len(table["rows"])
        table["column_count"] = len(table["headers"])
    return [table for table in tables if table["rows"]]


def _delimiter(text):
    # Metadata above a table defeats Sniffer. Score recognized headers per
    # delimiter before falling back to its ordinary dialect detection.
    candidates = []
    for delimiter in (",", ";", "\t", "|"):
        best = 0
        for row in csv.reader(io.StringIO(text[:262144]), delimiter=delimiter):
            matches = sum(bool(v) for v in header_candidates(row).values())
            best = max(best, matches if len(row) >= 2 else 0)
        candidates.append((best, delimiter))
    if max(score for score, _ in candidates) >= 2:
        return max(candidates, key=lambda item: item[0])[1]
    try:
        return csv.Sniffer().sniff(text[:65536], delimiters=",;\t|").delimiter
    except csv.Error:
        return ","


def _csv_tables(path):
    raw = path.read_bytes()
    warnings = []
    try:
        text = raw.decode("utf-16") if raw.startswith((b"\xff\xfe", b"\xfe\xff")) else raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("cp1252")
        warnings.append("File was decoded as Windows-1252; check names and symbols.")
    if "\x00" in text:
        raise ValueError("The CSV contains binary/null data or an unsupported encoding.")
    delimiter = _delimiter(text)
    records = []
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter, strict=True)
    try:
        while True:
            start_row = reader.line_num + 1
            try:
                row = next(reader)
            except StopIteration:
                break
            if reader.line_num > MAX_ROWS:
                raise ValueError(f"The file exceeds the {MAX_ROWS}-row import limit.")
            records.append((start_row, row, {}))
    except csv.Error as exc:
        raise ValueError("Malformed CSV; check quoting, delimiters, and field lengths.") from exc
    return _tables(records, initial_warnings=warnings)


def _xlsx_tables(path):
    from openpyxl import load_workbook

    with zipfile.ZipFile(path) as archive:
        if sum(entry.file_size for entry in archive.infolist()) > MAX_EXPANDED_BYTES:
            raise ValueError("Expanded workbook exceeds the supported size limit.")
        if len(archive.infolist()) > 10_000:
            raise ValueError("Workbook has too many components.")
    workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    cached = None
    total_rows = 0
    all_tables = []
    try:
        for sheet in workbook.worksheets:
            if (sheet.max_row or 0) > MAX_ROWS or (sheet.max_column or 0) > MAX_COLUMNS:
                raise ValueError("Worksheet dimensions exceed the supported row/column limit.")
            records = []
            sheet_warnings = []
            if sheet.sheet_state != "visible":
                sheet_warnings.append("This worksheet is hidden in the source workbook; review before importing.")
            cache_rows = None
            for row_number, cells in enumerate(sheet.iter_rows(), 1):
                total_rows += 1
                if total_rows > MAX_ROWS:
                    raise ValueError(f"Workbook exceeds the {MAX_ROWS}-row import limit.")
                values, formulas, warnings = [], {}, []
                for index, cell in enumerate(cells):
                    value = cell.value
                    if cell.data_type == "f":
                        formulas[index] = value
                        if cached is None:
                            cached = load_workbook(path, read_only=True, data_only=True, keep_links=False)
                        if cache_rows is None:
                            # Read once per worksheet, avoiding repeated XML scans.
                            cache_rows = list(cached[sheet.title].iter_rows(values_only=True))
                        value = cache_rows[row_number - 1][index]
                        if value is None:
                            value = formulas[index]
                            warnings.append("A formula has no cached value; relevant formulas must be replaced with values before import.")
                        else:
                            warnings.append("Cached formula value used; formulas were not executed. Check that the workbook was recalculated before export.")
                    values.append(value)
                records.append((row_number, values, {"formulas": formulas, "warnings": warnings}))
            all_tables.extend(_tables(records, sheet.title, sheet_warnings))
    finally:
        workbook.close()
        if cached:
            cached.close()
    return all_tables


def parse_file(path):
    """Extract table previews. Unsupported formats fail with an actionable error."""
    path = Path(path)
    if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
        raise ValueError("Structured portfolio import supports CSV, TSV, and XLSX. Export PDF, legacy XLS, or text statements as CSV/XLSX; scanned or unstructured statements cannot be safely calculated.")
    if path.stat().st_size > MAX_FILE_BYTES:
        raise ValueError(f"File exceeds the {MAX_FILE_BYTES // (1024 * 1024)} MB import limit.")
    try:
        tables = _xlsx_tables(path) if path.suffix.lower() == ".xlsx" else _csv_tables(path)
    except (zipfile.BadZipFile, KeyError) as exc:
        raise ValueError("Invalid or damaged XLSX workbook.") from exc
    if not tables:
        raise ValueError("No financial tables found. Use a header row with date, action, symbol, quantity, and price/amount columns, or review the export format.")
    for index, table in enumerate(tables):
        table["table_index"] = index
    return tables
